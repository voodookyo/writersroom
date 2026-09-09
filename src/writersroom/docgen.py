"""DOCX 生成与重开校验（docs/design/spec.md §4.8）。

- spec 提供 cover / sections（段落、无序/有序列表、表格）/ header / footer /
  font_cjk / 固定列宽；spec 之外不内置任何真实项目元数据
  （cover.title 缺失时只留空位，不编造默认值）；
- 依赖 python-docx（可选 extras）；页码用 oxml PAGE 域（fldChar/instrText）；
  中文字体写 w:eastAsia（run 级直写 + style 级）；
- 表格 autofit=False + tblLayout fixed + 逐列固定宽度；
- validate 重开文件校验：节数/段落数/表格数/每表列宽/页脚 PAGE 域 XML/eastAsia 字体，
  返回 {ok, checks: [{name, ok, detail}]}。
"""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

from .core import TOOL, WritersRoomError, now_iso, sha256_text

_COVER_ORDER = ["title", "subtitle", "author", "date"]  # 其余 cover 键按字典序追加
_COVER_SIZES = {"title": 26, "subtitle": 16}            # pt；其余字段 12pt
_CJK_STYLES = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
               "List Bullet", "List Number"]
_WIDTH_TOLERANCE_CM = 0.1  # 列宽校验容差（twips 换算取整引入的误差）


def _require_docx():
    try:
        import docx
        return docx
    except ImportError as e:
        raise WritersRoomError(
            "缺少 python-docx，无法生成 DOCX",
            "pip install 'writersroom[docx]' 或 pip install python-docx",
        ) from e


# ---------- 生成 ----------

def _set_cjk_on_rpr(rpr, font: str) -> None:
    from docx.oxml.ns import qn
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), font)


def _cjk_run(run, font: str | None) -> None:
    if font:
        _set_cjk_on_rpr(run._element.get_or_add_rPr(), font)


def _set_cjk_styles(document, font: str) -> None:
    for name in _CJK_STYLES:
        try:
            style = document.styles[name]
        except KeyError:
            continue
        _set_cjk_on_rpr(style.element.get_or_add_rPr(), font)


def _add_page_field(paragraph):
    """oxml 逃生舱：在段落中插入 PAGE 页码域（fldChar/instrText）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


def _render_cover(document, cover: dict, font: str | None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    for _ in range(4):  # 封面上方留白
        document.add_paragraph()
    keys = [k for k in _COVER_ORDER if k in cover]
    keys += sorted(k for k in cover if k not in _COVER_ORDER)
    if "title" not in cover:
        p = document.add_paragraph()  # 留空位，不编造默认标题
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k in keys:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(cover[k]))
        run.font.size = Pt(_COVER_SIZES.get(k, 12))
        run.font.bold = k == "title"
        _cjk_run(run, font)
    document.add_page_break()


def _render_header_footer(document, spec: dict, font: str | None) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    section = document.sections[0]
    header_text = spec.get("header")
    if header_text:
        run = section.header.paragraphs[0].add_run(str(header_text))
        _cjk_run(run, font)
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_text = spec.get("footer")
    if footer_text:
        _cjk_run(footer_para.add_run(str(footer_text) + " "), font)
    _cjk_run(_add_page_field(footer_para), font)  # 页码始终插入（结构元素，非元数据）


def _render_table(document, item: dict, font: str | None, usable_cm: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm
    headers = [str(h) for h in (item.get("headers") or [])]
    rows = [[str(c) for c in r] for r in (item.get("rows") or [])]
    ncols = len(headers) if headers else (len(rows[0]) if rows else 0)
    if ncols == 0:
        raise WritersRoomError("表格缺少列", "table 块需要 headers 或至少一行 rows")
    widths = item.get("col_widths_cm")
    if widths is None:
        widths = [round(usable_cm / ncols, 2)] * ncols  # 未指定时等分可用页宽
    if len(widths) != ncols:
        raise WritersRoomError(
            f"col_widths_cm 数量（{len(widths)}）与表格列数（{ncols}）不一致",
            "让 col_widths_cm 与 headers/rows 的列数一致",
        )
    table = document.add_table(rows=(1 if headers else 0) + len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    if tbl_pr.find(qn("w:tblLayout")) is None:  # 显式固定布局（双保险）
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)
    grid = list(table.rows)
    r0 = 0
    if headers:
        for j, text in enumerate(headers):
            run = grid[0].cells[j].paragraphs[0].add_run(text)
            run.font.bold = True
            _cjk_run(run, font)
        r0 = 1
    for i, row in enumerate(rows):
        for j in range(ncols):
            text = row[j] if j < len(row) else ""
            _cjk_run(grid[r0 + i].cells[j].paragraphs[0].add_run(text), font)
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(float(w))
        for cell in table.columns[j].cells:
            cell.width = Cm(float(w))


def _render_sections(document, sections: list[dict], font: str | None, usable_cm: float) -> None:
    for sec in sections:
        heading = document.add_heading(str(sec.get("heading", "")), level=int(sec.get("level", 1)))
        for run in heading.runs:
            _cjk_run(run, font)
        for item in sec.get("content", []):
            kind = item.get("type")
            if kind == "paragraph":
                _cjk_run(document.add_paragraph().add_run(str(item.get("text", ""))), font)
            elif kind == "list":
                style = "List Number" if item.get("ordered") else "List Bullet"
                for entry in item.get("items", []):
                    _cjk_run(document.add_paragraph(style=style).add_run(str(entry)), font)
            elif kind == "table":
                _render_table(document, item, font, usable_cm)
            else:
                raise WritersRoomError(
                    f"未知内容块类型：{kind}",
                    "sections[].content[].type 仅支持 paragraph / list / table",
                )


def generate(spec: dict, out: str | Path) -> dict:
    """按 spec 生成 DOCX；返回 {ok, out, counts, produced_by, created_at, source_hash}。"""
    _require_docx()
    from docx import Document
    if not isinstance(spec, dict):
        raise WritersRoomError("spec 必须是 JSON 对象", "检查 --spec 文件内容是否为合法 JSON 对象")
    document = Document()
    font = spec.get("font_cjk")
    font = str(font) if font else None
    if font:
        _set_cjk_styles(document, font)
    cover = spec.get("cover") or {}
    if cover:
        _render_cover(document, cover, font)
    _render_header_footer(document, spec, font)
    sec = document.sections[0]
    usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # 1cm = 360000 EMU
    _render_sections(document, spec.get("sections") or [], font, usable_cm)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    return {
        "ok": True,
        "out": str(out),
        "produced_by": TOOL,
        "created_at": now_iso(),
        "source_hash": sha256_text(json.dumps(spec, ensure_ascii=False, sort_keys=True)),
        "counts": {
            "sections": len(spec.get("sections") or []),
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
        },
    }


# ---------- 校验 ----------

def _part_xmls(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(
            z.read(n).decode("utf-8") for n in sorted(z.namelist())
            if n.startswith(prefix) and n.endswith(".xml")
        )


def validate(path: str | Path, spec: dict | None = None) -> dict:
    """重开 DOCX 校验结构；返回 {ok, checks: [{name, ok, detail}]}。

    提供 spec 时额外比对：表格数、每表固定列宽、页眉文本、封面标题、font_cjk 具体字体。
    """
    _require_docx()
    from docx import Document
    from docx.oxml.ns import qn

    checks: list[dict] = []

    def check(name: str, ok: bool, detail: str) -> None:
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    p = Path(path)
    check("file_exists", p.exists(), str(p))
    if not p.exists():
        return {"ok": False, "checks": checks}
    try:
        document = Document(str(p))
        check("opens", True, "python-docx 重开成功")
    except Exception as e:
        check("opens", False, f"重开失败：{e}")
        return {"ok": False, "checks": checks}

    check("sections", len(document.sections) >= 1, f"节数={len(document.sections)}")
    check("paragraphs", len(document.paragraphs) > 0, f"段落数={len(document.paragraphs)}")

    spec_tables: list[dict] = []
    if spec is not None:
        spec_tables = [it for s in spec.get("sections") or []
                       for it in s.get("content", []) if it.get("type") == "table"]
        check("tables", len(document.tables) == len(spec_tables),
              f"表格数={len(document.tables)}，期望={len(spec_tables)}")
    else:
        check("tables", True, f"表格数={len(document.tables)}")

    for idx, (table, tspec) in enumerate(zip(document.tables, spec_tables)):
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        fixed = layout is not None and layout.get(qn("w:type")) == "fixed"
        grid = table._tbl.find(qn("w:tblGrid"))
        actual = [round(int(gc.get(qn("w:w"))) * 2.54 / 1440, 2)
                  for gc in grid.findall(qn("w:gridCol"))]
        want = tspec.get("col_widths_cm")
        if want:
            width_ok = len(actual) == len(want) and all(
                abs(a - float(w)) < _WIDTH_TOLERANCE_CM for a, w in zip(actual, want))
            detail = f"fixed={fixed}，列宽cm={actual}，期望={want}"
        else:
            ncols = len(tspec.get("headers") or (tspec.get("rows") or [[]])[0])
            width_ok = len(actual) == ncols and all(a > 0 for a in actual)
            detail = f"fixed={fixed}，列宽cm={actual}（spec 未指定，等分）"
        check(f"table_{idx}_widths", fixed and width_ok, detail)

    footer_xml = _part_xmls(p, "word/footer")
    check("footer_page_field", "PAGE" in footer_xml and "instrText" in footer_xml,
          "页脚 XML 含 PAGE 域" if "PAGE" in footer_xml else "页脚缺 PAGE 域")

    if spec is not None and spec.get("header"):
        header_xml = _part_xmls(p, "word/header")
        check("header_text", str(spec["header"]) in header_xml,
              f"页眉含「{spec['header']}」" if str(spec["header"]) in header_xml else "页眉缺指定文本")

    styles_xml = _part_xmls(p, "word/styles")
    font = (spec or {}).get("font_cjk")
    if font:
        has = f'w:eastAsia="{font}"' in styles_xml
        check("eastasia_font", has,
              f"styles.xml 含 w:eastAsia={font}" if has else "styles.xml 缺 w:eastAsia 字体")
    else:
        check("eastasia_font", "eastAsia" in styles_xml,
              "styles.xml 含 eastAsia 声明" if "eastAsia" in styles_xml
              else "未设置 eastAsia（spec 未给 font_cjk 时允许）")

    if spec is not None and (spec.get("cover") or {}).get("title"):
        title = str(spec["cover"]["title"])
        found = any(par.text == title for par in document.paragraphs)
        check("cover_title", found, f"封面标题「{title}」" + ("存在" if found else "缺失"))

    return {"ok": all(c["ok"] for c in checks), "checks": checks}
